"""Инженерная проверка 20.10.2026: переход после удаления только внутри приложения; КП без номера заказа."""
import io
import re
from pathlib import Path

from docx import Document

from app.storage.orm import PositionRow
from tests.web.test_quick_web import calc, client_for, files


def test_delete_error_does_not_redirect_outside(app):
    client_for(app, "manager").post("/orders", files=files("D-1"))
    with app.state.session_factory() as s:
        pos_id = s.query(PositionRow).first().id
    c = client_for(app, "technolog")
    for back in ("https://evil.example/x", "//evil.example/x", "/\\evil.example"):
        resp = c.post(f"/positions/{pos_id}/delete", data={"reason": "", "back": back}, follow_redirects=False)
        assert resp.status_code == 303 and resp.headers["location"] == f"/positions/{pos_id}", back


def test_quick_proposal_has_no_empty_order_line(app):
    c = client_for(app)
    _, token = calc(c)
    page = c.post(f"/quick/{token}/proposal")
    assert page.status_code == 200 and "Коммерческое предложение" in page.text
    assert "к заказу" not in page.text
    kp_id = re.search(r"/proposals/(\d+)/file", page.text).group(1)
    doc = Document(io.BytesIO(c.get(f"/proposals/{kp_id}/file").content))
    text = "\n".join(p.text for p in doc.paragraphs) + "".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert "к заказу" not in text


def test_current_quantity_break_is_highlighted():
    css = (Path(__file__).resolve().parents[2] / "app" / "web" / "static" / "style.css").read_text(encoding="utf-8")
    assert ".quick-break-current" in css
