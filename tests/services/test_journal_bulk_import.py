"""Загрузка журнала калибровки одним действием: Excel (детали, время операций, замеры) + STEP/PDF."""
from openpyxl import load_workbook

from app.parsers.journal_excel import JOURNAL_SHEET, export_journal, import_journal, write_journal_template
from app.services import journal_service as js
from app.services import reference_service as ref
from app.services.file_check import check_entry_files
from app.services.order_service import UploadedFile
from tests.domain.conftest import make_entry
from tests.fixtures.samples import make_box_step, make_pdf
from tests.services.conftest import TECH


def files(stem):
    return [UploadedFile(f"{stem}.stp", make_box_step(100, 60, 20).encode()),
            UploadedFile(f"{stem}.pdf", make_pdf([stem.split()[0], "Material: D16T"]))]


def test_template_has_hint_and_reads_back_empty(tmp_path):
    path = write_journal_template(tmp_path / "t.xlsx")
    wb = load_workbook(path)
    assert {JOURNAL_SHEET, "Замеры", "Подсказка"} <= set(wb.sheetnames)
    assert import_journal(path) == []


def test_import_with_files_attaches_by_designation(session, tmp_path):
    xlsx = tmp_path / "j.xlsx"
    export_journal([make_entry("KPU.100.001"), make_entry("KPU.100.002")], xlsx)
    upload = files("KPU.100.001 Планка") + files("KPU.100.002") + files("KPU.999.999 Лишняя")
    report = js.import_journal_with_files(session, TECH, xlsx, upload, tmp_path / "jf")
    assert report.imported == 2 and report.attached == 4
    assert any("KPU.999.999" in w for w in report.warnings)
    for code in ("KPU.100.001", "KPU.100.002"):
        assert all(s.ok for s in check_entry_files(js.entry_row_by_code(session, code)))
    assert len(ref.journal_entries(session)) == 11  # 9 демо + 2 новые


def test_reimport_updates_times_and_keeps_files(session, tmp_path):
    xlsx = tmp_path / "j.xlsx"
    export_journal([make_entry("KPU.100.001")], xlsx)
    js.import_journal_with_files(session, TECH, xlsx, files("KPU.100.001"), tmp_path / "jf")
    entry = make_entry("KPU.100.001")
    entry.stages[3] = entry.stages[3].__class__(entry.stages[3].stage, 0.9, entry.stages[3].count_mode)
    export_journal([entry], xlsx)
    js.import_journal_with_files(session, TECH, xlsx, [], tmp_path / "jf")
    row = js.entry_row_by_code(session, "KPU.100.001")
    assert next(s for s in row.stages if s["stage"] == "machining")["hours"] == 0.9
    assert row.step_path and row.pdf_path  # файлы прежней загрузки на месте
