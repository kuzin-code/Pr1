"""Коммерческое предложение (простой вариант): заказчик, позиции, цена операций, срок ±10 %, желаемая дата."""
from datetime import date

import pytest

from app.domain.access import AccessDenied
from app.services import intake_service as intake
from app.services import proposal_service as kp
from tests.services.conftest import MASTER, MON, OPERATOR, part_files
from tests.services.test_quick_intake import MANAGER


@pytest.fixture
def order(session, tmp_path):
    return intake.quick_intake(session, MANAGER, part_files("Кронштейн-1") + part_files("Планка-2"),
                               tmp_path / "up", customer="ООО «Альфа»", desired_date=date(2026, 12, 1),
                               start_date=MON, quantity=4)


def test_proposal_from_order(session, order):
    p = kp.build_proposal(session, MANAGER, order.id)
    assert p.customer == "ООО «Альфа»" and p.number == "КП-" + order.number.split("-")[-1]
    assert [i.quantity for i in p.items] == [4, 4]
    first = p.items[0]
    assert first.machine_hours > 0 and first.machine_rate == 2000 and first.bench_rate == 1000
    assert first.price == pytest.approx(first.batch_price / 4)
    assert first.batch_price == pytest.approx(first.machine_hours * 2000 + first.bench_hours * 1000)
    assert p.total == pytest.approx(sum(i.batch_price for i in p.items))
    assert p.vat == pytest.approx(p.total * 0.2) and p.grand_total == pytest.approx(p.total * 1.2)
    assert p.range.earliest <= p.calc_date.date() <= p.range.latest
    assert p.desired == date(2026, 12, 1) and p.desired_check.label


def test_position_without_calc_has_no_price(session, tmp_path):
    files = part_files("Кронштейн-1") + [part_files("Планка-2")[1]]
    order = intake.quick_intake(session, MASTER, files, tmp_path / "up", start_date=MON)
    p = kp.build_proposal(session, MASTER, order.id)
    assert [i.price is None for i in p.items] == [False, True]
    assert "не рассчитана" in " ".join(p.notes)


def test_word_file_is_a4(session, order, tmp_path):
    from docx import Document
    from docx.shared import Mm
    path = kp.proposal_docx(session, MANAGER, order.id, tmp_path / "kp.docx")
    doc = Document(path)
    assert abs(doc.sections[0].page_width - Mm(210)) < Mm(1) and abs(doc.sections[0].page_height - Mm(297)) < Mm(1)
    text = "\n".join(p.text for p in doc.paragraphs) + "".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Коммерческое предложение" in text and "ООО «Альфа»" in text and "Кронштейн-1" in text
    assert "Станочная обработка" in text


def test_operator_cannot_make_proposal(session, order):
    with pytest.raises(AccessDenied):
        kp.build_proposal(session, OPERATOR, order.id)


def test_proposal_recorded_in_history(session, order):
    from app.storage.orm import HistoryRow
    kp.build_proposal(session, MANAGER, order.id, record_history=True)
    assert session.query(HistoryRow).filter_by(field="proposal").count() == 1
