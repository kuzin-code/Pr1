"""Ступени количества в быстром расчёте и КП: цена и срок для 1 / 10 / 50 шт., рабочая база не меняется."""
from datetime import date

import pytest
from docx import Document

from app.domain.access import AccessDenied
from app.services import intake_service as intake
from app.services import proposal_history as history
from app.services import quick_service as quick
from app.services.actor import ServiceError
from app.storage.orm import ProposalRow
from tests.services.conftest import MON, OPERATOR, part_files
from tests.services.test_quick_calc import counts
from tests.services.test_quick_intake import MANAGER


@pytest.fixture
def result(session, tmp_path):
    return quick.calculate(session, MANAGER, part_files("Кронштейн-1") + part_files("Планка-2"), tmp_path / "work",
                           quantity=10, start_date=MON, quantities=[50, 1, 10])


def test_breaks_for_every_position_without_touching_database(session, tmp_path):
    before = counts(session)
    s = quick.calculate(session, MANAGER, part_files(), tmp_path / "work", quantity=10, start_date=MON,
                        quantities=[1, 10, 50]).summary
    assert counts(session) == before
    assert s["quantities"] == [1, 10, 50] and s["params"]["quantities"] == [1, 10, 50]
    (item,) = s["items"]
    assert [b["quantity"] for b in item["breaks"]] == [1, 10, 50]
    prices = [b["piece_price"] for b in item["breaks"]]
    assert prices[0] > prices[-1] and all(a >= b for a, b in zip(prices, prices[1:]))
    assert all(b["ready"] for b in item["breaks"])
    assert item["breaks"][0]["ready"] <= item["breaks"][-1]["ready"]


def test_break_equal_to_quantity_matches_item_price(result):
    for item in result.summary["items"]:
        same = next(b for b in item["breaks"] if b["quantity"] == item["quantity"])
        assert same["batch_price"] == pytest.approx(item["batch_price"])
        assert same["machine_hours"] == pytest.approx(item["machine_hours"])


def test_no_breaks_by_default(session, tmp_path):
    s = quick.calculate(session, MANAGER, part_files(), tmp_path / "work", start_date=MON).summary
    assert s["quantities"] == [] and s["items"][0]["breaks"] == []


def test_change_breaks_and_recalculate(session, result, tmp_path):
    before = counts(session)
    s = quick.set_quantities(session, MANAGER, tmp_path / "work", result.token, "1, 100")
    assert s["quantities"] == [1, 100] and [b["quantity"] for b in s["items"][0]["breaks"]] == [1, 100]
    assert quick.load(tmp_path / "work", result.token)["quantities"] == [1, 100]  # сохранено
    assert counts(session) == before
    with pytest.raises(ServiceError):
        quick.set_quantities(session, MANAGER, tmp_path / "work", result.token, "1, 2, 3, 4, 5")
    assert quick.load(tmp_path / "work", result.token)["quantities"] == [1, 100]
    with pytest.raises(AccessDenied):
        quick.set_quantities(session, OPERATOR, tmp_path / "work", result.token, "1")


def test_breaks_survive_position_edit(session, result, tmp_path):
    d = result.summary["positions"][0]["designation"]
    s = quick.update_position(session, MANAGER, tmp_path / "work", result.token, d, {"quantity": 3})
    assert s["quantities"] == [1, 10, 50] and len(s["items"][0]["breaks"]) == 3


def _docx_text(path) -> str:
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] +
                     [c.text for t in doc.tables for r in t.rows for c in r.cells])


def test_proposal_from_quick_keeps_all_breaks(session, result, tmp_path):
    row = history.create_from_quick(session, MANAGER, tmp_path / "work", result.token, tmp_path / "kp")
    p = history.proposal_of(row)
    assert [b["quantity"] for b in p.items[0].breaks] == [1, 10, 50]
    assert p.items[0].breaks == result.summary["items"][0]["breaks"]
    assert p.grand_total == pytest.approx(result.summary["grand_total"])  # итог — по количеству позиции
    text = _docx_text(row.file_path)
    assert "Цена в зависимости от количества" in text and "50" in text


def test_proposal_from_order_with_breaks(session, tmp_path):
    order = intake.quick_intake(session, MANAGER, part_files("Кронштейн-1"), tmp_path / "up", start_date=MON,
                                desired_date=date(2026, 12, 1), quantity=4)
    plain = history.proposal_of(history.create_from_order(session, MANAGER, order.id, tmp_path / "kp"))
    assert plain.items[0].breaks == []
    row = history.create_from_order(session, MANAGER, order.id, tmp_path / "kp", quantities=[1, 4, 20])
    item = history.proposal_of(row).items[0]
    assert [b["quantity"] for b in item.breaks] == [1, 4, 20] and all(b["ready"] for b in item.breaks)
    assert item.breaks[1]["batch_price"] == pytest.approx(item.batch_price)
    assert "Цена в зависимости от количества" in _docx_text(row.file_path)


def test_old_proposal_without_breaks_opens(session, tmp_path):
    order = intake.quick_intake(session, MANAGER, part_files("Кронштейн-1"), tmp_path / "up", start_date=MON)
    row = history.create_from_order(session, MANAGER, order.id, tmp_path / "kp")
    data = dict(row.data)
    data["items"] = [{k: v for k, v in i.items() if k != "breaks"} for i in data["items"]]
    old = ProposalRow(number="КП-старое", created_at=row.created_at, created_by="manager", source="order",
                      customer="", designations=row.designations, grand_total=row.grand_total, data=data)
    session.add(old)
    session.commit()
    p = history.proposal_of(session.get(ProposalRow, old.id))
    assert p.items[0].breaks == [] and p.grand_total == pytest.approx(row.grand_total)
