"""Выгрузка в 1С: заказы за период, сумма КП, плановая и фактическая себестоимость, история, права."""
import csv
import io
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from app.domain.access import AccessDenied
from app.domain.accounting_export import headers
from app.services import export_service as export
from app.services import intake_service as intake
from app.services import operation_service as ops
from app.services import proposal_history as history
from app.services.history_service import list_history
from tests.services.conftest import HEAD, MASTER, MON, TECH, part_files
from tests.services.test_operation_flow import CUTTER, KONTROL, OPERATOR, SLESAR, approved
from tests.services.test_quick_intake import MANAGER

TODAY = date.today()


def _rows(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig")), delimiter=";"))


def test_order_with_proposal_and_plan(session, tmp_path):
    order = intake.quick_intake(session, MANAGER, part_files("Кронштейн-1") + part_files("Планка-2"),
                                tmp_path / "up", customer="ООО «Альфа»", start_date=MON, quantity=4)
    kp = history.create_from_order(session, MANAGER, order.id, tmp_path / "kp")
    name, kind, data = export.export_1c(session, MANAGER, TODAY, TODAY)
    assert name.endswith(".csv") and kind.startswith("text/csv")
    rows = _rows(data)
    assert rows[0] == headers() and len(rows) == 2
    line = dict(zip(headers(), rows[1]))
    assert line["Номер заказа"] == order.number and line["Заказчик"] == "ООО «Альфа»"
    assert line["Позиции"] == "Кронштейн-1; Планка-2" and line["Количество, шт."] == "8"
    assert line["Сумма КП, руб."] == f"{kp.grand_total:.2f}".replace(".", ",")
    assert line["Плановая себестоимость, руб."] and line["Фактическая себестоимость, руб."] == ""
    assert line["Готовность"] == "по плану" and line["Дата готовности"]


def test_actual_cost_and_ready_date_for_done_order(session, make_order):
    pos = approved(session, make_order)
    rows = ops.position_operations(session, pos.id)
    moment = datetime(2026, 10, 12, 9)
    for actor, row in zip((CUTTER, OPERATOR, OPERATOR, SLESAR, SLESAR, KONTROL), rows):
        ops.start(session, actor, row.id, moment)
        moment += timedelta(hours=1)
        ops.finish(session, actor, row.id, moment, reason="tool")
    line = dict(zip(headers(), _rows(export.export_1c(session, HEAD, TODAY, TODAY)[2])[1]))
    assert line["Фактическая себестоимость, руб."] not in ("", "0,00")
    assert (line["Готовность"], line["Дата готовности"]) == ("фактически", "12.10.2026")


def test_empty_period_header_only_and_history(session, make_order):
    make_order()
    name, _, data = export.export_1c(session, HEAD, date(2020, 1, 1), date(2020, 1, 31))
    assert _rows(data) == [headers()] and name == "zakazy_1c_20200101-20200131.csv"
    h = [x for x in list_history(session, "export") if x.field == "export_1c"][0]
    assert h.new == {"from": "2020-01-01", "to": "2020-01-31", "orders": 0, "format": "csv"}
    assert h.username == "head"


def test_xlsx(session, make_order):
    make_order()
    name, kind, data = export.export_1c(session, HEAD, TODAY, TODAY, "xlsx")
    sheet = load_workbook(io.BytesIO(data)).active
    assert name.endswith(".xlsx") and "spreadsheetml" in kind
    assert [c.value for c in sheet[1]] == headers() and sheet.max_row == 2


def test_rights_and_bad_input(session):
    for actor in (TECH, MASTER, OPERATOR):
        with pytest.raises(AccessDenied):
            export.export_1c(session, actor, TODAY, TODAY)
    with pytest.raises(ValueError, match="раньше"):
        export.export_1c(session, HEAD, TODAY, TODAY - timedelta(days=1))
    with pytest.raises(ValueError, match="Формат"):
        export.export_1c(session, HEAD, TODAY, TODAY, "pdf")
