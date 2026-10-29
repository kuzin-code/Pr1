"""ФТ-25, Д-4: загрузка журнала калибровки из Excel."""
import pytest
from openpyxl import Workbook

from app.domain.enums import CountMode, Stage, StageGroup
from app.parsers.journal_excel import (
    JOURNAL_SHEET,
    JournalFormatError,
    export_journal,
    import_journal,
    journal_headers,
)
from tests.domain.conftest import machine_measurements, make_entry


def test_roundtrip_nine_parts(tmp_path):
    journal = [make_entry(f"P-{i}", dims=(100 + i, 60, 20), measurements=machine_measurements(1.1))
               for i in range(9)]
    path = tmp_path / "journal.xlsx"
    export_journal(journal, path)
    loaded = import_journal(path)
    assert len(loaded) == 9
    first = loaded[0]
    assert first.part_code == "P-0"
    assert first.dims.as_list() == [100, 60, 20]
    assert first.stages == journal[0].stages
    assert first.measurements[0].group is StageGroup.MACHINE
    assert first.measurements[0].ratio == pytest.approx(1.1)


def test_empty_stage_hours_mean_stage_absent(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = JOURNAL_SHEET
    headers = journal_headers()
    ws.append(headers)
    row = {h: None for h in headers}
    row.update({"Обозначение": "X-1", "Наименование": "Втулка", "Станок": "ST20Y", "Материал": "Сталь 45",
                "Габарит 1, мм": 40, "Габарит 2, мм": 40, "Габарит 3, мм": 60, "Квалитет": 7,
                "Балл сложности": 2, "Технология": "turning",
                "Обработка на станке, ч": 0.4, "Обработка на станке, счёт": "штука",
                "Наладка, ч": 1, "Наладка, счёт": "партия"})
    ws.append([row[h] for h in headers])
    path = tmp_path / "j.xlsx"
    wb.save(path)
    (entry,) = import_journal(path)
    assert [s.stage for s in entry.stages] == [Stage.SETUP, Stage.MACHINING]
    assert entry.stages[1].count_mode is CountMode.PER_PIECE
    assert entry.technology == "turning"


def test_missing_column_error(tmp_path):
    wb = Workbook()
    wb.active.title = JOURNAL_SHEET
    wb.active.append(["Обозначение"])
    path = tmp_path / "j.xlsx"
    wb.save(path)
    with pytest.raises(JournalFormatError, match="Нет столбцов"):
        import_journal(path)


def test_bad_row_reports_row_number(tmp_path):
    path = tmp_path / "j.xlsx"
    export_journal([make_entry()], path)
    from openpyxl import load_workbook
    wb = load_workbook(path)
    wb[JOURNAL_SHEET]["E2"] = "много"
    wb.save(path)
    with pytest.raises(JournalFormatError, match="строка 2"):
        import_journal(path)


def _sheet_row(tmp_path, **values):
    wb = Workbook()
    ws = wb.active
    ws.title = JOURNAL_SHEET
    headers = journal_headers()
    ws.append(headers)
    row = {h: None for h in headers}
    row.update({"Обозначение": "XWS.65.A.1.P001040 Оправа зеркала", "Наименование": "Оправа", "Материал": "Д16Т",
                "Габарит 1, мм": 90, "Габарит 2, мм": 50, "Габарит 3, мм": 20, "Квалитет": 12,
                "Балл сложности": 2, "Технология": "milling", "Наладка, счёт": "партия",
                "Обработка на станке, счёт": "штука", "Слесарная, ч": 0.25} | values)
    ws.append([row.get(h) for h in headers])
    path = tmp_path / "j.xlsx"
    wb.save(path)
    return path


def test_second_machine_through_slash(tmp_path):
    """«VOLZH1/SMM», «0,6/0,7», «0,3/0,2»: до черты — основной станок, после — второй; время не складывается."""
    path = _sheet_row(tmp_path, **{"Станок": "VOLZH1/SMM", "Наладка, ч": "0,6/0,7", "Обработка на станке, ч": "0,3/0,2"})
    (entry,) = import_journal(path)
    assert entry.machine_code == "VOLZH1"
    hours = {s.stage: s.hours for s in entry.stages}
    assert hours[Stage.SETUP] == 0.6 and hours[Stage.MACHINING] == 0.3
    assert entry.extra_machines == [{"machine": "SMM", "setup": 0.7, "machining": 0.2}]


def test_row_without_slash_unchanged(tmp_path):
    path = _sheet_row(tmp_path, **{"Станок": "VOLZH1", "Наладка, ч": 0.6, "Обработка на станке, ч": 0.3})
    (entry,) = import_journal(path)
    assert entry.machine_code == "VOLZH1" and entry.extra_machines == []


def test_second_machine_without_time_rejected(tmp_path):
    path = _sheet_row(tmp_path, **{"Станок": "VOLZH1/SMM", "Наладка, ч": 0.6, "Обработка на станке, ч": 0.3})
    with pytest.raises(JournalFormatError, match="второго станка"):
        import_journal(path)


def test_export_second_machine_through_slash(tmp_path):
    entry = make_entry("P-2", machine="VOLZH1")
    entry.extra_machines = [{"machine": "SMM", "setup": 0.7, "machining": 0.2}]
    path = tmp_path / "j.xlsx"
    export_journal([entry], path)
    from openpyxl import load_workbook
    ws = load_workbook(path)[JOURNAL_SHEET]
    values = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert values["Станок"] == "VOLZH1/SMM"
    assert values["Наладка, ч"] == "1,5/0,7" and values["Обработка на станке, ч"] == "0,5/0,2"
    (back,) = import_journal(path)
    assert back.extra_machines == entry.extra_machines and back.stages == entry.stages
