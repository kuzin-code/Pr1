"""Список деталей заказа по шаблону предприятия (ФТ-32)."""
from datetime import date

import pytest

from app.parsers.order_list import OrderListError, cooperation_hint, read_order_list
from tests.fixtures.samples import make_order_list

ITEMS = [
    ("MBIS.AI.11000.20", "Лист обшивки нижний", "1", "Altec X0 Ral7047 (4 мм/0,4 мм)"),
    ("MBIS.AI.11000.25", "Обшивка крыши задняя", 18, "Altec X0 Ral7047 (4 мм/0,4 мм)", "", "срочно"),
    ("MBIS.AI.18000.01", "Основание", 12, "Д16Т ГОСТ 4784-2019", "Анодирование чёрное"),
]


def test_header_and_items(tmp_path):
    data = read_order_list(make_order_list(tmp_path / "z.xlsx", ITEMS, sections={1: "Детали из алюмокомпозита",
                                                                               3: "Детали из алюминия"}))
    assert data.customer == "ООО «Тест»"
    assert data.order_date == date(2026, 3, 12)
    assert data.contacts["Разработчик"] == "Иванов И. И."
    assert data.contacts["Тел (ответственный)"] == "+7 900 000-00-02"
    assert [i.designation for i in data.items] == ["MBIS.AI.11000.20", "MBIS.AI.11000.25", "MBIS.AI.18000.01"]
    first, second, third = data.items
    assert first.quantity == 1 and second.quantity == 18  # количество текстом и числом
    assert first.section == "Детали из алюмокомпозита" and third.section == "Детали из алюминия"
    assert second.note == "срочно" and third.coating == "Анодирование чёрное"
    assert third.material == "Д16Т ГОСТ 4784-2019"


def test_other_sheets_ignored(tmp_path):
    data = read_order_list(make_order_list(tmp_path / "z.xlsx", ITEMS[:1], extra_sheet=True))
    assert len(data.items) == 1


@pytest.mark.parametrize("bad,match", [
    ([("A-1", "Деталь", "0", "Д16Т")], "строка 8.*количество"),
    ([("A-1", "Деталь", "много", "Д16Т")], "строка 8.*количество"),
    ([("A-1", "Деталь", 1, "Д16Т"), ("A-1", "Деталь", 2, "Д16Т")], "повторяется"),
])
def test_row_errors(tmp_path, bad, match):
    with pytest.raises(OrderListError, match=match):
        read_order_list(make_order_list(tmp_path / "z.xlsx", bad))


def test_no_table_error(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)
    with pytest.raises(OrderListError, match="Обозначение"):
        read_order_list(path)


@pytest.mark.parametrize("text,hint", [
    ("Анодирование чёрное", True), ("Окраска RAL 9005", True), ("гальваника Ц9", True),
    ("термообработка HRC 40", True), ("", False), ("срочно", False), (None, False),
])
def test_cooperation_hint(text, hint):
    assert bool(cooperation_hint(text)) is hint


def test_generated_template_is_readable_after_filling(tmp_path):
    from openpyxl import load_workbook
    from app.parsers.order_list import write_template
    path = write_template(tmp_path / "Шаблон.xlsx")
    wb = load_workbook(path)
    ws = wb["Заказ"]
    assert "Подсказка" in wb.sheetnames
    row = next(r for r in range(1, 30) if ws.cell(r, 3).value == "Обозначение")
    ws.cell(row + 1, 2, 1)
    ws.cell(row + 1, 3, "KPU.001")
    ws.cell(row + 1, 4, "Корпус")
    ws.cell(row + 1, 5, 3)
    ws["D2"] = "ООО «Заказчик»"
    wb.save(path)
    data = read_order_list(path)
    assert [(i.designation, i.quantity) for i in data.items] == [("KPU.001", 3)]
    assert data.customer == "ООО «Заказчик»"
