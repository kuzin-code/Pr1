"""Импорт/экспорт журнала калибровки в Excel (ФТ-25, Д-4).

Лист «Журнал»: одна строка — одна деталь, по каждому этапу два столбца: «<этап>, ч» и «<этап>, счёт».
Пустое время — этапа нет в маршруте. Лист «Замеры»: факт по партии и группе этапов.
"""
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.domain.enums import CountMode, Stage, StageGroup
from app.domain.models import CalibrationEntry, Dimensions, Measurement, StageNorm

JOURNAL_SHEET = "Журнал"
MEASUREMENTS_SHEET = "Замеры"
BASE_COLUMNS = ["Обозначение", "Наименование", "Станок", "Материал", "Габарит 1, мм", "Габарит 2, мм",
                "Габарит 3, мм", "Квалитет", "Балл сложности", "Технология"]
MEASUREMENT_COLUMNS = ["Обозначение", "Группа этапов", "Расчёт, ч", "Факт, ч"]
COUNT_LABELS = {"штука": CountMode.PER_PIECE, "партия": CountMode.PER_BATCH, "разово": CountMode.ONCE}
GROUP_LABELS = {"станочная": StageGroup.MACHINE, "слесарно-финишная": StageGroup.BENCH}


class JournalFormatError(ValueError):
    pass


def journal_headers() -> list[str]:
    headers = list(BASE_COLUMNS)
    for stage in Stage:
        headers += [f"{stage.label}, ч", f"{stage.label}, счёт"]
    return headers


def _label(mapping: dict, value) -> str:
    return next(k for k, v in mapping.items() if v is value)


def export_journal(journal: list[CalibrationEntry], path: Path | str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = JOURNAL_SHEET
    ws.append(journal_headers())
    ms = wb.create_sheet(MEASUREMENTS_SHEET)
    ms.append(MEASUREMENT_COLUMNS)
    for e in journal:
        row = [e.part_code, e.name, e.machine_code, e.material, *e.dims.as_list(), e.tolerance_grade,
               e.complexity, e.technology]
        stages = {s.stage: s for s in e.stages}
        for stage in Stage:
            norm = stages.get(stage)
            row += [norm.hours, _label(COUNT_LABELS, norm.count_mode)] if norm else [None, None]
        ws.append(row)
        for m in e.measurements:
            ms.append([e.part_code, _label(GROUP_LABELS, m.group), m.calc_hours, m.actual_hours])
    wb.save(path)


def _num(value, row: int, column: str, cast=float):
    try:
        return cast(value)
    except (TypeError, ValueError):
        raise JournalFormatError(f"{JOURNAL_SHEET}, строка {row}: некорректное значение «{column}»") from None


def _read_entry(values: dict, row: int) -> CalibrationEntry:
    stages = []
    for stage in Stage:
        hours = values.get(f"{stage.label}, ч")
        if hours in (None, ""):
            continue
        mode_label = str(values.get(f"{stage.label}, счёт") or "штука").strip().lower()
        if mode_label not in COUNT_LABELS:
            raise JournalFormatError(f"{JOURNAL_SHEET}, строка {row}: неизвестный способ счёта «{mode_label}»")
        stages.append(StageNorm(stage, _num(hours, row, stage.label), COUNT_LABELS[mode_label]))
    try:
        dims = Dimensions(*(_num(values[c], row, c) for c in BASE_COLUMNS[4:7]))
    except ValueError as err:
        raise JournalFormatError(f"{JOURNAL_SHEET}, строка {row}: {err}") from None
    return CalibrationEntry(
        part_code=str(values["Обозначение"]).strip(),
        name=str(values.get("Наименование") or "").strip(),
        machine_code=str(values.get("Станок") or "").strip(),
        material=str(values.get("Материал") or "").strip(),
        dims=dims,
        tolerance_grade=_num(values["Квалитет"], row, "Квалитет", int),
        complexity=_num(values["Балл сложности"], row, "Балл сложности", int),
        stages=stages,
        technology=str(values.get("Технология") or "milling").strip(),
    )


def _rows(ws, required: list[str]):
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    missing = [c for c in required if c not in headers]
    if missing:
        raise JournalFormatError(f"Лист «{ws.title}»: Нет столбцов: {', '.join(missing)}")
    for idx, values in enumerate(rows, start=2):
        if any(v not in (None, "") for v in values):
            yield idx, dict(zip(headers, values))


def import_journal(path: Path | str) -> list[CalibrationEntry]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as err:  # noqa: BLE001
        raise JournalFormatError(f"Файл Excel не читается: {err}") from None
    if JOURNAL_SHEET not in wb.sheetnames:
        raise JournalFormatError(f"Нет листа «{JOURNAL_SHEET}»")
    entries = {}
    for row, values in _rows(wb[JOURNAL_SHEET], BASE_COLUMNS):
        entry = _read_entry(values, row)
        entries[entry.part_code] = entry
    if MEASUREMENTS_SHEET in wb.sheetnames:
        for row, values in _rows(wb[MEASUREMENTS_SHEET], MEASUREMENT_COLUMNS):
            code = str(values["Обозначение"]).strip()
            group = GROUP_LABELS.get(str(values["Группа этапов"]).strip().lower())
            if code not in entries or group is None:
                raise JournalFormatError(f"{MEASUREMENTS_SHEET}, строка {row}: неизвестная деталь или группа")
            entries[code].measurements.append(Measurement(
                group, _num(values["Расчёт, ч"], row, "Расчёт, ч"), _num(values["Факт, ч"], row, "Факт, ч")))
    wb.close()
    return list(entries.values())


JOURNAL_HINTS = [
    "Как заполнить журнал калибровки",
    "1. Лист «Журнал»: одна строка — одна деталь. Обозначение — как в основной надписи чертежа (графа 2).",
    "2. Станок — код из справочника: VF2, SMM, MV6, VOLZH1, ST20Y. Технология — milling (фрезерная) или turning (токарная).",
    "3. По каждой операции — время в часах и способ счёта: «штука», «партия» или «разово». "
    "Пустое время — операции нет в маршруте. Наладка всегда считается на партию.",
    "4. Лист «Замеры» (необязательно): фактическое время прошлых партий по группе «станочная» или "
    "«слесарно-финишная» — расчётное и фактическое на партию, ч. Из них система считает коэффициент k.",
    "5. Чертежи PDF и модели STEP загружаются вместе с Excel: имя файла должно начинаться с обозначения "
    "детали (например, «КПУ.741645.001 Корпус.pdf»).",
    "6. Повторная загрузка обновляет время у деталей с тем же обозначением; уже прикреплённые файлы сохраняются.",
]


def write_journal_template(path: Path | str) -> Path:
    """Пустой журнал калибровки с подсказкой по заполнению."""
    export_journal([], path)
    wb = load_workbook(path)
    hints = wb.create_sheet("Подсказка")
    for i, line in enumerate(JOURNAL_HINTS, start=1):
        hints.cell(i, 1, line)
    hints.column_dimensions["A"].width = 130
    wb.save(path)
    return Path(path)
