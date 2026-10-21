"""Список деталей заказа по шаблону предприятия «Шаблон заказа» (ФТ-32).

Лист с таблицей «№ | Обозначение | Наименование | К-во | Материал | Покрытие, цвет | Прим.»;
выше таблицы — шапка (дата заказа, заказчик, разработчик, ответственный, телефоны, e-mail);
строки без обозначения внутри таблицы — названия разделов («Детали из алюмокомпозита»).
"""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from app.domain.order_intake import MAX_QUANTITY, QUANTITY_LIMIT_MESSAGE

COLUMNS = {
    "designation": ("обозначение",),
    "name": ("наименование",),
    "quantity": ("к-во", "кол-во", "количество"),
    "material": ("материал",),
    "coating": ("покрытие",),
    "note": ("прим",),
}
HEADER_LABELS = ("дата заказа", "желаем", "срок", "заказчик", "разработчик", "ответственный", "тел", "e-mail", "email")
# Внешние (кооперационные) операции по покрытию и примечаниям (ТЗ, раздел 10).
COOPERATION = re.compile(r"анод|оксид|гальван|цинк|хромир|никел|кадми|окрас|покрас|порошк|эмал|\bral\b|лак"
                         r"|термообр|закалк|отпуск|цементац|азотир|шлифов|лазерн", re.I)


class OrderListError(ValueError):
    pass


@dataclass
class OrderListItem:
    row: int
    designation: str
    name: str
    quantity: int
    material: str = ""
    coating: str = ""
    note: str = ""
    section: str = ""


@dataclass
class OrderList:
    order_date: date | None = None
    desired_date: date | None = None  # желаемая дата изготовления (срок заказчика)
    customer: str = ""
    contacts: dict[str, str] = field(default_factory=dict)
    items: list[OrderListItem] = field(default_factory=list)


def cooperation_hint(text: str | None) -> str:
    """Найденное указание на внешнюю операцию (покрытие, термообработка…) или пустая строка."""
    found = COOPERATION.search(text or "")
    return found.group(0) if found else ""


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def _find_table(ws):
    for row in ws.iter_rows():
        columns = {}
        for cell in row:
            label = _text(cell.value).casefold()
            for key, names in COLUMNS.items():
                if key not in columns and any(label.startswith(n) for n in names):
                    columns[key] = cell.column
        if {"designation", "quantity"} <= set(columns):
            return row[0].row, columns
    return None, None


def _read_header(ws, header_row: int, data: OrderList) -> None:
    for row in ws.iter_rows(max_row=header_row - 1):
        cells = [c for c in row if _text(c.value)]
        pair = 0
        for i, cell in enumerate(cells):
            label = _text(cell.value).rstrip(":").strip()
            key = label.casefold()
            if not any(key.startswith(h) for h in HEADER_LABELS) or i + 1 >= len(cells):
                continue
            value = cells[i + 1].value
            if any(_text(value).casefold().startswith(h) for h in HEADER_LABELS):
                pair += 1
                continue
            day = value.date() if isinstance(value, datetime) else value if isinstance(value, date) else None
            if key.startswith(("желаем", "срок")):
                data.desired_date = day
            elif key.startswith("дата"):
                data.order_date = day
            elif key.startswith("заказчик"):
                data.customer = _text(value)
            elif key.startswith(("разработчик", "ответственный")):
                data.contacts[label] = _text(value)
            else:
                data.contacts[f"{label} ({'разработчик' if pair == 0 else 'ответственный'})"] = _text(value)
            pair += 1


def _quantity(value, row: int) -> int:
    try:
        number = float(str(value).replace(",", ".").strip())
    except (TypeError, ValueError):
        raise OrderListError(f"Список деталей, строка {row}: количество «{_text(value)}» — нужно целое число") from None
    if number < 1 or not number.is_integer():
        raise OrderListError(f"Список деталей, строка {row}: количество «{_text(value)}» — нужно целое число ≥ 1")
    if number > MAX_QUANTITY:
        raise OrderListError(f"Список деталей, строка {row}: количество {int(number)} — "
                             f"{QUANTITY_LIMIT_MESSAGE[0].lower()}{QUANTITY_LIMIT_MESSAGE[1:]}")
    return int(number)


def read_order_list(path: Path | str) -> OrderList:
    try:
        wb = load_workbook(path, data_only=True)
    except Exception as err:  # noqa: BLE001
        raise OrderListError(f"Файл списка не читается: {err}") from None
    for ws in wb.worksheets:
        header_row, columns = _find_table(ws)
        if header_row:
            break
    else:
        raise OrderListError("В списке не найдена таблица со столбцами «Обозначение» и «К-во»")

    data = OrderList()
    _read_header(ws, header_row, data)
    section, seen = "", set()
    for row in ws.iter_rows(min_row=header_row + 1):
        values = {key: row[col - 1].value if col - 1 < len(row) else None for key, col in columns.items()}
        designation = _text(values.get("designation"))
        if not designation:
            texts = [_text(c.value) for c in row if _text(c.value)]
            if len(texts) == 1 and values.get("quantity") in (None, ""):
                section = texts[0]  # строка раздела
            continue
        if designation.casefold() in seen:
            raise OrderListError(f"Список деталей, строка {row[0].row}: обозначение «{designation}» повторяется")
        seen.add(designation.casefold())
        data.items.append(OrderListItem(
            row=row[0].row, designation=designation, name=_text(values.get("name")),
            quantity=_quantity(values.get("quantity"), row[0].row), material=_text(values.get("material")),
            coating=_text(values.get("coating")), note=_text(values.get("note")), section=section))
    if not data.items:
        raise OrderListError("В списке нет деталей")
    return data


TEMPLATE_HINTS = [
    "Как заполнить список деталей",
    "1. Шапка: заказчик, дата заказа, разработчик и ответственный с контактами (необязательно).",
    "2. Таблица: одна строка — одна деталь. Обязательны «Обозначение» и «К-во» "
    f"(целое число от 1 до {MAX_QUANTITY}).",
    "3. Раздел (например, «Детали из алюминия») — строка, где заполнена только первая ячейка.",
    "4. «Покрытие, цвет» и «Прим.»: анодирование, окраска, RAL, гальваника, термообработка, шлифовка "
    "отмечаются системой как внешняя (кооперационная) операция — расчёт «без кооперации».",
    "5. Материал из списка используется, если его не удалось прочитать из чертежа; главный документ — чертёж PDF.",
    "6. Файлы STEP и PDF загружаются вместе со списком; имя файла должно начинаться с обозначения детали, "
    "например «MBIS.AI.18000.01 Основание.step». Детали без файлов можно дополнить позже в карточке позиции.",
    "7. В заказе — не более 50 деталей. Столбцы можно переставлять: система ищет их по заголовкам.",
]


def write_template(path: Path | str, rows: int = 50) -> Path:
    """Чистый шаблон списка деталей на основе шаблона заказа предприятия."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Заказ"
    thin, medium = Side(style="thin"), Side(style="medium")
    ws.merge_cells("B1:H1")
    ws["B1"] = "Заказ на изготовление деталей"
    ws["B1"].font = Font(bold=True, size=14)
    header = [("Заказчик:", "Дата заказа:"), ("Разработчик:", "Ответственный:"), ("Тел:", "Тел:"),
              ("E-mail:", "E-mail:")]
    for r, (left, right) in enumerate(header, start=2):
        ws.merge_cells(f"B{r}:C{r}")
        ws.merge_cells(f"D{r}:E{r}")
        ws.merge_cells(f"G{r}:H{r}")
        ws[f"B{r}"], ws[f"F{r}"] = left, right
        for cell in (f"D{r}", f"G{r}"):
            ws[cell].border = Border(bottom=thin)
    ws["G2"].number_format = "DD.MM.YYYY"
    ws.merge_cells("B6:C6")
    ws.merge_cells("D6:E6")
    ws["B6"] = "Желаемая дата изготовления:"
    ws["D6"].number_format = "DD.MM.YYYY"
    ws["D6"].border = Border(bottom=thin)
    titles = ["№", "Обозначение", "Наименование", "К-во", "Материал", "Покрытие, цвет", "Прим."]
    widths = [5, 28, 40, 7, 40, 26, 40]
    fill = PatternFill("solid", fgColor="DCE6F1")
    for col, (title, width) in enumerate(zip(titles, widths), start=2):
        cell = ws.cell(7, col, title)
        cell.font, cell.fill = Font(bold=True), fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=medium, bottom=medium, left=thin, right=thin)
        ws.column_dimensions[cell.column_letter].width = width
    for r in range(8, 8 + rows):
        for col in range(2, 9):
            ws.cell(r, col).border = Border(left=thin, right=thin, bottom=thin)
    quantity = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MAX_QUANTITY),
                              allow_blank=True, error=f"Количество — целое число от 1 до {MAX_QUANTITY}",
                              showErrorMessage=True)
    ws.add_data_validation(quantity)
    quantity.add(f"E8:E{7 + rows}")
    ws.freeze_panes = "B8"
    hints = wb.create_sheet("Подсказка")
    for i, line in enumerate(TEMPLATE_HINTS, start=1):
        hints.cell(i, 1, line).font = Font(bold=i == 1)
    hints.column_dimensions["A"].width = 120
    wb.save(path)
    return Path(path)
