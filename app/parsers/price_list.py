"""Прайс заготовок prutki.ru (price.xls или та же таблица, сохранённая как .xlsx).

Шапка — строки до заголовка таблицы «Код | Товар | Цена | Остатки (шт)»; в ней дата «Цены и наличие от 26.10.2026»
и ставка «Включают НДС 22%». Строка товара: «Заготовка плита алюминиевая Д16Т 50 х 100 х 200 мм» —
форма, прилагательное, марка, размеры через «х». Аноды, строки без цены и нераспознанные строки пропускаются.
"""
import io
import re
from datetime import datetime

from app.domain.blank_price import PriceItem, PriceList

FORMS = ("лист", "плита", "пруток", "круг", "стержень", "шестигранник", "квадрат", "полоса", "труба", "уголок")
_NUM = r"(\d+(?:[.,]\d+)?)"
# размеры — числа через « х » с пробелами: «40Х 20 х 100» — марка 40Х, Ø20, длина 100
_ITEM = re.compile(rf"^Заготовка\s+(\S+)\s+(\S+)\s+(?:(.*?)\s+)??({_NUM}(?:\s+[хx×]\s+{_NUM})*)\s*мм$", re.I)
_DATE = re.compile(r"от\s+(\d{2}\.\d{2}\.\d{4})")
_VAT = re.compile(r"НДС\s*(\d+(?:[.,]\d+)?)\s*%", re.I)
HEADER_ROWS = 15  # заголовок таблицы ищется в первых строках


class PriceFormatError(ValueError):
    pass


def _rows(content: bytes, filename: str) -> list[list]:
    name = (filename or "").casefold()
    try:
        if name.endswith(".xls"):
            import xlrd

            sheet = xlrd.open_workbook(file_contents=content).sheet_by_index(0)
            return [sheet.row_values(r) for r in range(sheet.nrows)]
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                return [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]
            finally:
                book.close()
    except Exception:  # повреждённый файл или другой формат под тем же расширением
        raise PriceFormatError("Файл не читается как таблица Excel: скачайте price.xls с сайта заново") from None
    raise PriceFormatError("Прайс — файл Excel: price.xls с сайта (или .xlsx)")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _number(value) -> float | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    try:
        return float(_text(value).replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def _code(value) -> str:
    number = _number(value)
    return f"{number:g}" if number is not None and number == int(number) else _text(value)


def parse_price_list(content: bytes, filename: str) -> tuple[PriceList, int]:
    """Прайс и число пропущенных строк. Не похоже на прайс prutki.ru — PriceFormatError с понятной причиной."""
    rows = _rows(content, filename)
    head = next((n for n, r in enumerate(rows[:HEADER_ROWS])
                 if [_text(c).casefold() for c in r[:3]] == ["код", "товар", "цена"]), None)
    if head is None:
        raise PriceFormatError("Не найден заголовок таблицы «Код | Товар | Цена | Остатки»: это не прайс prutki.ru "
                               "или изменился формат файла")
    header = " ".join(_text(c) for r in rows[:head] for c in r)
    date = _DATE.search(header)
    vat = _VAT.search(header)
    items, skipped = [], 0
    for row in rows[head + 1:]:
        row = list(row) + [None] * 4
        match = _ITEM.match(_text(row[1]))
        price = _number(row[2])
        if not match or match.group(1).casefold() not in FORMS or not price or price <= 0:
            skipped += 1 if any(_text(c) for c in row[:4]) else 0
            continue
        dims = tuple(float(x.replace(",", ".")) for x in re.findall(_NUM, match.group(4)))
        items.append(PriceItem(_code(row[0]), _text(row[1]), match.group(1).casefold(), match.group(3) or "", dims,
                               price, _number(row[3])))
    if not items:
        raise PriceFormatError("В файле нет ни одной строки «Заготовка … мм» с ценой: проверьте, что это прайс "
                               "prutki.ru")
    when = datetime.strptime(date.group(1), "%d.%m.%Y").date().isoformat() if date else None
    return PriceList(when, float(vat.group(1).replace(",", ".")) if vat else None, tuple(items)), skipped
