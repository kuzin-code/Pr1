"""Коммерческое предложение в Word, лист А4 (поля 20 / 15 / 15 / 15 мм): текст правится в Word."""
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from app.services.proposal_service import Proposal

ACCENT = "993556"
MUTED = RGBColor(0x66, 0x70, 0x85)
CHECK_COLORS = {"ok": RGBColor(0x1E, 0x7A, 0x46), "range": RGBColor(0xB5, 0x47, 0x08),
                "late": RGBColor(0xB4, 0x23, 0x18)}
RIGHT = WD_ALIGN_PARAGRAPH.RIGHT


def money(value: float | None) -> str:
    if value is None:
        return "—"
    return f"{value:,.2f}".replace(",", " ").replace(".", ",").replace(",00", "")


def hours(value: float) -> str:
    return f"{value:.2f}".rstrip("0").rstrip(".").replace(".", ",")


def ready(value: str | None) -> str:
    return datetime.fromisoformat(value).strftime("%d.%m.%Y") if value else "—"


def _shade(cell, color: str) -> None:
    shd = OxmlElement("w:shd")
    for key, value in (("val", "clear"), ("color", "auto"), ("fill", color)):
        shd.set(qn(f"w:{key}"), value)
    cell._tc.get_or_add_tcPr().append(shd)


def _border(par, color: str = ACCENT, size: int = 18) -> None:
    bdr, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", str(size)), ("space", "4"), ("color", color)):
        bottom.set(qn(f"w:{key}"), value)
    bdr.append(bottom)
    par._p.get_or_add_pPr().append(bdr)


def _widths(table, widths_mm) -> None:
    """Фиксированные ширины колонок: Word и LibreOffice читают tblGrid и tcW."""
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table._tbl.tblPr.append(layout)
    for grid, width in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths_mm):
        grid.set(qn("w:w"), str(int(Mm(width).twips)))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_mm):
            cell.width = Mm(width)


def _run(par, text, bold=False, size=None, color=None):
    r = par.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def _cell(cell, text, bold=False, align=None, size=9.5, color=None):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align:
        p.alignment = align
    _run(p, text, bold, size, color)


def _table(doc, heads, rows, widths, right_from):
    t = doc.add_table(rows=1, cols=len(heads))
    t.style, t.alignment, t.autofit = "Table Grid", WD_TABLE_ALIGNMENT.CENTER, False
    for c, head in enumerate(heads):
        _cell(t.rows[0].cells[c], head, bold=True, size=9, align=RIGHT if c >= right_from else None)
        _shade(t.rows[0].cells[c], "F4F5F7")
    for values in rows:
        cells = t.add_row().cells
        for c, value in enumerate(values):
            _cell(cells[c], value, align=RIGHT if c >= right_from else None)
    return t


def _breaks(doc, p: Proposal) -> None:
    """Ступени количества: цена за штуку и за партию, готовность на станке для 1 / 10 / 50 шт."""
    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(8)
    _run(sub, "Цена в зависимости от количества", bold=True)
    rows = []
    for i, it in enumerate(p.items, start=1):
        for n, b in enumerate(it.breaks):
            rows.append([str(i) if n == 0 else "", it.designation if n == 0 else "", str(b["quantity"]),
                         money(b["piece_price"]), money(b["batch_price"]), ready(b.get("ready"))])
    widths = [8, 52, 22, 28, 30, 35]
    _widths(_table(doc, ["№", "Обозначение", "Кол-во, шт.", "Цена за шт., ₽", "Сумма за партию, ₽",
                         "Готовность на станке"], rows, widths, right_from=2), widths)
    _run(doc.add_paragraph(), p.breaks_note, size=8.5, color=MUTED)


def write_docx(p: Proposal, path: Path) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Mm(15), Mm(15), Mm(20), Mm(15)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    head = doc.add_table(rows=1, cols=2)
    left, right = head.rows[0].cells
    _cell(left, p.org, bold=True, size=12)
    _run(left.add_paragraph(), p.requisites, size=9, color=MUTED)
    for i, line in enumerate((f"КП № {p.number} от {p.date:%d.%m.%Y}", f"к заказу {p.order_number}",
                              f"Действительно до {p.valid_until:%d.%m.%Y}")):
        par = right.paragraphs[0] if i == 0 else right.add_paragraph()
        par.alignment, par.paragraph_format.space_after = RIGHT, Pt(0)
        _run(par, line, size=9, color=MUTED)
    _widths(head, [115, 60])
    _border(doc.add_paragraph())

    title = doc.add_paragraph()
    title.paragraph_format.space_before, title.paragraph_format.space_after = Pt(8), Pt(6)
    _run(title, "Коммерческое предложение", bold=True, size=17)
    par = doc.add_paragraph()
    _run(par, "Заказчик: ", color=MUTED), _run(par, p.customer or "—", bold=True)
    if p.contact_line:
        par = doc.add_paragraph()
        _run(par, "Контакты: ", color=MUTED), _run(par, p.contact_line)
    doc.add_paragraph("Благодарим за запрос. Предлагаем изготовить детали по вашей конструкторской документации "
                      "на следующих условиях:")

    rows = [[str(i), it.designation, it.name, it.material, str(it.quantity), money(it.price), money(it.batch_price)]
            for i, it in enumerate(p.items, start=1)]
    t = _table(doc, ["№", "Обозначение", "Наименование", "Материал", "Кол-во, шт.", "Цена за шт., ₽", "Сумма, ₽"],
               rows, [8, 42, 34, 26, 17, 23, 25], right_from=4)
    for label, value in (("Итого без НДС", p.total), (f"НДС {p.vat_percent} %", p.vat), ("Итого с НДС", p.grand_total)):
        cells = t.add_row().cells
        merged = cells[0].merge(cells[5])
        _cell(merged, label, bold=True, align=RIGHT)
        _cell(cells[6], money(value), bold=True, align=RIGHT)
    _widths(t, [8, 42, 34, 26, 17, 23, 25])

    if p.has_breaks:
        _breaks(doc, p)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_before = Pt(8)
    _run(sub, "Цена операций (на партию)", bold=True)
    ops = []
    for i, it in enumerate(p.items, start=1):
        if it.batch_price is None:
            ops.append([str(i), it.designation, "—", "—", "цена не рассчитана"])
            continue
        ops.append([str(i), it.designation, "Станочная обработка (наладка и обработка)",
                    f"{hours(it.machine_hours)} ч × {money(it.machine_rate)}", money(it.machine_price)])
        ops.append(["", "", "Слесарно-финишная обработка", f"{hours(it.bench_hours)} ч × {money(it.bench_rate)}",
                    money(it.bench_price)])
    _widths(_table(doc, ["№", "Обозначение", "Операция", "Время × ставка, ₽/ч", "Сумма, ₽"], ops,
                   [8, 42, 60, 40, 25], right_from=3), [8, 42, 60, 40, 25])
    _run(doc.add_paragraph(), p.price_note, size=8.5, color=MUTED)

    terms = doc.add_table(rows=0, cols=2)
    terms.style = "Table Grid"

    def term(label, parts, note=""):
        a, b = terms.add_row().cells
        _cell(a, label, color=MUTED, size=10)
        _shade(a, "F8F9FB")
        par = b.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        for text, bold, size, color in parts:
            _run(par, text, bold, size, color)
        if note:
            _run(b.add_paragraph(), note, size=8.5, color=MUTED)

    if p.range:
        term("Расчётная дата изготовления", [(f"{p.range.finish:%d.%m.%Y}", True, 10, None),
                                             ("  — по текущей загрузке станков с ЧПУ", False, 9, MUTED)])
        term("Срок изготовления", [(f"с {p.range.earliest:%d.%m.%Y} по {p.range.latest:%d.%m.%Y}", True, 12, None)],
             f"±10 % от длительности изготовления ({p.range.duration} раб. дн.), округлено вверх: "
             f"±{p.range.delta} раб. дн.")
    else:
        term("Срок изготовления", [("уточняется", True, 10, None)])
    if p.desired:
        term("Желаемая дата заказчика", [(f"{p.desired:%d.%m.%Y}", True, 10, None),
                                         (f"  {p.desired_check.label}", True, 10, CHECK_COLORS.get(p.desired_check.status))])
    term("Начало изготовления", [(f"{p.start:%d.%m.%Y}; после согласования КП и предоплаты, при более позднем "
                                  "подтверждении срок пересчитывается", False, 10, None)])
    term("Условия оплаты", [(p.payment, False, 10, None)])
    term("Доставка", [(p.delivery, False, 10, None)])
    _widths(terms, [55, 120])

    for note in p.notes:
        _run(doc.add_paragraph(), note, size=9, color=CHECK_COLORS["late"])
    doc.add_paragraph()
    _run(doc.add_paragraph(), p.disclaimer, size=9, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    signs = doc.add_table(rows=1, cols=2)
    for cell, text in zip(signs.rows[0].cells, ("Менеджер производства / Ф. И. О. / подпись",
                                                "Начальник участка / Ф. И. О. / подпись")):
        _border(cell.paragraphs[0], color="1D2330", size=6)
        _run(cell.add_paragraph(), text, size=9, color=MUTED)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
